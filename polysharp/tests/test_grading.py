import json
import os
import time

import pytest

from polysharp.grading import (Top10Log, build_xlsx, grade_pending, grade_row, history_breakdown,
                               league_label, log_breakdown, market_type, settle)
from polysharp.store import Store

NOW = int(time.time())


def gm(cid, outs, prices, tokens, closed=True):
    return {"conditionId": cid, "closed": closed, "outcomes": json.dumps(outs),
            "outcomePrices": json.dumps([str(p) for p in prices]), "clobTokenIds": json.dumps(tokens)}


def test_labels():
    assert league_label("nfl-pit-cle-2026-10-02") == "NFL" and league_label("will-trump-win") == "OTHER"
    assert market_type("Spread: Browns (-3.5)") == "Spread"
    assert market_type("Eagles vs. Titans: O/U 39.5") == "Total"
    assert market_type("Dodgers vs. Padres") == "Moneyline"
    assert market_type("Exact Score: Germany 1 - 1 Serbia?") == "Prop"


def test_settle_real_gamma_shape():
    # Orioles/Rockies O/U 11.5 as Gamma returns it: Under won
    m = gm("0x8a", ["Over", "Under"], [0, 1], ["9595", "9920"])
    assert settle(m, "9595", "Over") == "L" and settle(m, "9920", "Under") == "W"
    assert settle(m, "zzz", "Under") == "W"                       # falls back to outcome name
    assert settle(gm("0x1", ["A", "B"], [0.5, 0.5], ["1", "2"]), "1", "A") == "P"
    assert settle(gm("0x1", ["A", "B"], [0.6, 0.4], ["1", "2"], closed=False), "1", "A") is None


def test_grade_row_units():
    row = {"cost": 7000, "shares": 10000, "avg": 0.7}
    assert grade_row(row, "W") == (3000, pytest.approx(0.4286, abs=1e-4))
    assert grade_row(row, "L") == (-7000, -1.0)


def test_history_breakdown_groups_and_aligns_window():
    closed = [{"asset": "a", "eventSlug": "nfl-x", "title": "Spread: Browns (-3.5)", "totalBought": 1000,
               "avgPrice": 0.5, "realizedPnl": 500, "timestamp": NOW - 86400},
              {"asset": "b", "eventSlug": "nfl-y", "title": "Eagles vs. Titans: O/U 39.5", "totalBought": 1000,
               "avgPrice": 0.5, "realizedPnl": -500, "timestamp": NOW - 3600},
              {"asset": "c", "eventSlug": "atp-z", "title": "Japan Open: A vs B", "totalBought": 2000,
               "avgPrice": 0.5, "realizedPnl": 900, "timestamp": NOW - 7200}]
    from datetime import datetime
    from zoneinfo import ZoneInfo
    iso = lambda t: datetime.fromtimestamp(t, ZoneInfo("UTC")).isoformat()  # noqa: E731
    dead = [{"asset": "d", "eventSlug": "mlb-q", "title": "Dodgers vs. Padres", "curPrice": 0,
             "initialValue": 400, "endDate": iso(NOW - 3600)},                 # in window: counts
            {"asset": "e", "eventSlug": "mlb-q", "title": "Dodgers vs. Padres", "curPrice": 0,
             "initialValue": 9999, "endDate": iso(NOW - 90 * 86400)}]          # ancient: ignored
    h = history_breakdown(closed, dead)
    lg = {r["key"]: r for r in h["leagues"]}
    assert h["n"] == 4 and set(lg) == {"NFL", "ATP", "MLB"}
    assert lg["ATP"]["pnl"] == 900 and lg["ATP"]["roi"] == pytest.approx(0.9)
    assert lg["NFL"]["n"] == 2 and lg["NFL"]["pnl"] == 0 and lg["NFL"]["win_rate"] == 0.5
    assert lg["MLB"]["pnl"] == -400
    ty = {r["key"]: r for r in h["types"]}
    assert ty["Spread"]["pnl"] == 500 and ty["Total"]["pnl"] == -500


class GammaAPI:
    def __init__(self, markets):
        self.markets = markets

    async def gamma_markets(self, cids):
        return [self.markets[c] for c in cids if c in self.markets]


def pos(cid, asset, title, outcome, cost, avg, slug):
    return {"conditionId": cid, "asset": asset, "title": title, "outcome": outcome,
            "initialValue": cost, "size": cost / avg, "avgPrice": avg, "eventSlug": slug}


@pytest.mark.asyncio
async def test_log_snapshot_grade_and_report(tmp_path):
    st = Store(str(tmp_path / "g.db"))
    lg = Top10Log(st)
    meta = {"c1": {"league": "nfl", "game_start": NOW - 5 * 3600},
            "c2": {"league": "mlb", "game_start": NOW - 5 * 3600},
            "c3": {"league": "nfl", "game_start": NOW + 5 * 3600}}            # not started: not graded
    rows = [pos("c1", "t1", "Spread: Browns (-3.5)", "Steelers", 7000, 0.7, "nfl-pit-cle"),
            pos("c2", "t2", "Dodgers vs. Padres", "Dodgers", 5000, 0.5, "mlb-lad-sd"),
            pos("c3", "t3", "Bills vs. Jets", "Bills", 3000, 0.6, "nfl-buf-nyj")]
    assert lg.record("0xa", "alwaysfade", rows, meta) == 3
    # same position again next day: still one row, size refreshed
    rows[0]["initialValue"] = 8000
    lg.record("0xa", "alwaysfade", rows[:1], meta, now=NOW + 86400)
    assert len(lg.rows()) == 3 and lg.rows()[0]["cost"] == 8000
    api = GammaAPI({"c1": gm("c1", ["Steelers", "Browns"], [1, 0], ["t1", "x"]),
                    "c2": gm("c2", ["Dodgers", "Padres"], [0, 1], ["t2", "y"])})
    assert await grade_pending(api, lg) == 2
    by = {r["pick"]: r for r in lg.rows()}
    assert by["Steelers +3.5"]["result"] == "W" and by["Dodgers"]["result"] == "L"
    assert by["Bills"]["result"] is None
    rep = log_breakdown(lg.rows("0xa"))
    lgs = {r["key"]: r for r in rep["leagues"]}
    assert lgs["NFL"]["w"] == 1 and lgs["MLB"]["l"] == 1 and lgs["MLB"]["units"] == -1.0

    # Excel export opens and has the three sheets + live formulas
    from openpyxl import load_workbook
    path = build_xlsx(str(tmp_path / "x.xlsx"), lg.rows(), ["alwaysfade"])
    wb = load_workbook(path)
    assert wb.sheetnames == ["Plays", "By League", "By Type"]
    assert wb["Plays"].max_row == 4 and wb["Plays"]["J4"].value == "open"
    assert wb["By League"]["C2"].value.startswith("=COUNTIFS(")
    assert {wb["By League"][f"B{i}"].value for i in (2, 3)} == {"MLB", "NFL"}
