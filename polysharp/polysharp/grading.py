"""Where is each account actually good? Two views, both grouped by league and bet type.

1. HISTORY (instant): every settled position Polymarket has for the wallet
   (latest `SPORTS_HISTORY` settled bets + unredeemed losers from that same window).

2. LOG (builds over time): once a day the bot snapshots each feed wallet's top-10
   open sports positions into `top10_log`. An hourly grader reads the final result
   from Gamma (`outcomePrices` ["1","0"] etc.) and records W/L/push, their $ P&L and
   a flat-stake result in units (win = (1-avg)/avg u, loss = -1u) so accounts
   with very different bet sizes compare fairly.
"""
import asyncio
import json
import logging
import time
from collections import defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo

from .markets import LEAGUES, parse_ts

log = logging.getLogger(__name__)

SCHEMA = """
CREATE TABLE IF NOT EXISTS top10_log (
    wallet TEXT, name TEXT, asset TEXT, condition_id TEXT, day TEXT, logged_ts REAL,
    title TEXT, outcome TEXT, pick TEXT, league TEXT, mtype TEXT, event_slug TEXT,
    cost REAL, shares REAL, avg REAL, game_start REAL,
    result TEXT, pnl REAL, units REAL, graded_ts REAL,
    PRIMARY KEY (wallet, asset));
"""


# ----------------------------------------------------------------- labels
def league_label(slug):
    p = (slug or "").lower().split("-", 1)[0]
    return p.upper() if p in LEAGUES else "OTHER"


def market_type(title):
    t = (title or "").lower()
    if t.startswith("spread:") or "spread" in t:
        return "Spread"
    if "o/u" in t or "over/under" in t or "total" in t:
        return "Total"
    if "exact score" in t or "draw" in t or "both teams" in t or "first inning" in t \
            or "run scored" in t:
        return "Prop"
    if " vs" in t or "will " in t and " win" in t:
        return "Moneyline"
    return "Other"


def _empty():
    return {"n": 0, "wins": 0, "cost": 0.0, "pnl": 0.0}


def _add(b, cost, pnl):
    b["n"] += 1
    b["cost"] += cost
    b["pnl"] += pnl
    b["wins"] += pnl > 0


def _finish(groups):
    out = []
    for k, b in groups.items():
        out.append({"key": k, **b, "roi": b["pnl"] / b["cost"] if b["cost"] else 0.0,
                    "win_rate": b["wins"] / b["n"] if b["n"] else 0.0})
    out.sort(key=lambda r: -r["pnl"])
    return out


# ----------------------------------------------------------------- history
def history_breakdown(closed, dead, now=None):
    """Group settled positions by league and by bet type. Same-window rule as win_sample."""
    now = now or time.time()
    by_league, by_type = defaultdict(_empty), defaultdict(_empty)
    if not closed:
        return {"leagues": [], "types": [], "n": 0, "days": 0}
    start = min(int(p.get("timestamp") or now) for p in closed)
    seen = {str(p.get("asset")) for p in closed}
    rows = []
    for p in closed:
        cost = float(p.get("totalBought") or 0) * float(p.get("avgPrice") or 0)
        if cost > 0:
            rows.append((p, cost, float(p.get("realizedPnl") or 0)))
    for p in dead:
        if float(p.get("curPrice") or 0) > 0 or str(p.get("asset")) in seen:
            continue
        end = parse_ts(p.get("endDate"))
        cost = float(p.get("initialValue") or 0)
        if end and end >= start - 86400 and cost > 0:
            rows.append((p, cost, -cost))
    for p, cost, pnl in rows:
        _add(by_league[league_label(p.get("eventSlug") or p.get("slug"))], cost, pnl)
        _add(by_type[market_type(p.get("title"))], cost, pnl)
    return {"leagues": _finish(by_league), "types": _finish(by_type), "n": len(rows),
            "days": round((now - start) / 86400, 1)}


# ----------------------------------------------------------------- daily log
class Top10Log:
    def __init__(self, store):
        self.db = store.db
        self.db.executescript(SCHEMA)
        self.db.commit()

    def record(self, wallet, name, rows, meta, tz="America/Halifax", now=None):
        """Upsert today's top-10 rows. A position already logged keeps its first day but
        its size/avg is refreshed until it's graded."""
        from .watcher import pick_label
        now = now or time.time()
        day = datetime.fromtimestamp(now, ZoneInfo(tz)).strftime("%Y-%m-%d")
        n = 0
        for r in rows:
            m = meta.get(r.get("conditionId")) or {}
            slug = r.get("eventSlug") or r.get("slug") or ""
            self.db.execute(
                """INSERT INTO top10_log (wallet,name,asset,condition_id,day,logged_ts,title,outcome,
                       pick,league,mtype,event_slug,cost,shares,avg,game_start)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(wallet, asset) DO UPDATE SET
                       cost=excluded.cost, shares=excluded.shares, avg=excluded.avg, name=excluded.name
                   WHERE top10_log.result IS NULL""",
                (wallet, name, str(r.get("asset")), r.get("conditionId"), day, now,
                 r.get("title"), r.get("outcome"), pick_label(r.get("title"), r.get("outcome")),
                 (m.get("league") or "").upper() or league_label(slug), market_type(r.get("title")),
                 slug, float(r.get("initialValue") or 0), float(r.get("size") or 0),
                 float(r.get("avgPrice") or 0), m.get("game_start")))
            n += 1
        self.db.commit()
        return n

    def pending(self, now=None, limit=400):
        now = now or time.time()
        rows = self.db.execute(
            """SELECT wallet, asset, condition_id, outcome, cost, shares, avg FROM top10_log
               WHERE result IS NULL AND (game_start IS NULL OR game_start < ?)
               ORDER BY logged_ts LIMIT ?""", (now - 2 * 3600, limit)).fetchall()
        return [dict(r) for r in rows]

    def set_result(self, wallet, asset, result, pnl, units):
        self.db.execute("""UPDATE top10_log SET result=?, pnl=?, units=?, graded_ts=?
                           WHERE wallet=? AND asset=?""",
                        (result, pnl, units, time.time(), wallet, asset))

    def rows(self, wallet=None):
        q, args = "SELECT * FROM top10_log", ()
        if wallet:
            q, args = q + " WHERE wallet=?", (wallet,)
        return [dict(r) for r in self.db.execute(q + " ORDER BY logged_ts", args).fetchall()]


def settle(market, asset, outcome):
    """('W'|'L'|'P'|None) from a Gamma market row. None = not resolved yet."""
    if not market or not market.get("closed"):
        return None
    try:
        outs = json.loads(market.get("outcomes") or "[]")
        prices = [float(x) for x in json.loads(market.get("outcomePrices") or "[]")]
        tokens = [str(x) for x in json.loads(market.get("clobTokenIds") or "[]")]
    except (ValueError, TypeError):
        return None
    if not prices:
        return None
    idx = tokens.index(str(asset)) if str(asset) in tokens else \
        (outs.index(outcome) if outcome in outs else None)
    if idx is None or idx >= len(prices):
        return None
    p = prices[idx]
    if p >= 0.99:
        return "W"
    if p <= 0.01:
        return "L"
    if 0.4 <= p <= 0.6 and max(prices) < 0.99:
        return "P"           # voided / 50-50 settlement
    return None


def grade_row(row, result):
    cost, shares, avg = row["cost"], row["shares"], row["avg"] or 0.5
    if result == "W":
        return shares - cost, (1 - avg) / avg
    if result == "L":
        return -cost, -1.0
    return shares * 0.5 - cost, 0.5 / avg - 1


async def grade_pending(api, log10):
    todo = log10.pending()
    if not todo:
        return 0
    cids = sorted({r["condition_id"] for r in todo if r["condition_id"]})
    markets = {}
    for i in range(0, len(cids), 20):
        try:
            for m in await api.gamma_markets(cids[i:i + 20]):
                markets[m.get("conditionId")] = m
        except Exception as e:
            log.debug("grade lookup failed: %s", e)
    n = 0
    for r in todo:
        res = settle(markets.get(r["condition_id"]), r["asset"], r["outcome"])
        if res:
            pnl, units = grade_row(r, res)
            log10.set_result(r["wallet"], r["asset"], res, round(pnl, 2), round(units, 4))
            n += 1
    log10.db.commit()
    return n


def log_breakdown(rows):
    """League / type tables from graded log rows, in flat-stake units and $."""
    by_league, by_type = defaultdict(lambda: {"n": 0, "w": 0, "l": 0, "units": 0.0, "pnl": 0.0}), \
        defaultdict(lambda: {"n": 0, "w": 0, "l": 0, "units": 0.0, "pnl": 0.0})
    for r in rows:
        if not r.get("result"):
            continue
        for b in (by_league[r["league"] or "OTHER"], by_type[r["mtype"] or "Other"]):
            b["n"] += 1
            b["w"] += r["result"] == "W"
            b["l"] += r["result"] == "L"
            b["units"] += r["units"] or 0
            b["pnl"] += r["pnl"] or 0

    def fin(d):
        return sorted(({"key": k, **v} for k, v in d.items()), key=lambda x: -x["units"])
    return {"leagues": fin(by_league), "types": fin(by_type)}


# ----------------------------------------------------------------- excel
def build_xlsx(path, rows, names):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, c in enumerate(ws[1], 1):
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Plays"
    header(ws, ["Day", "Account", "League", "Type", "Market", "Pick", "Avg", "Cost $",
                "Shares", "Result", "P&L $", "Units"],
           [11, 18, 8, 10, 48, 26, 7, 11, 11, 7, 11, 8])
    for r in rows:
        ws.append([r["day"], r["name"], r["league"], r["mtype"], r["title"], r["pick"],
                   r["avg"], round(r["cost"], 2), round(r["shares"], 2), r["result"] or "open",
                   r["pnl"], r["units"]])
    n = ws.max_row
    for col, fmt in (("G", "0.000"), ("H", "#,##0"), ("I", "#,##0"), ("K", "#,##0;[Red]-#,##0"),
                     ("L", "0.00;[Red]-0.00")):
        for c in ws[f"{col}2:{col}{max(n, 2)}"]:
            c[0].number_format = fmt
    ws.auto_filter.ref = f"A1:L{max(n, 1)}"

    # Pivot sheets: live COUNTIFS/SUMIFS over Plays (one row per account x key seen)
    def pivot(title, label, key_col, key_field):
        p = wb.create_sheet(title)
        header(p, ["Account", label, "Graded", "W", "L", "Win %", "Units", "P&L $", "Units/bet"],
               [18, 10, 8, 6, 6, 8, 9, 12, 10])
        n2 = max(n, 2)
        rng, krng, res = f"Plays!$B$2:$B${n2}", f"Plays!${key_col}$2:${key_col}${n2}", f"Plays!$J$2:$J${n2}"
        combos = sorted({(r["name"], r[key_field]) for r in rows}, key=lambda x: (x[0] or "", x[1] or ""))
        for i, (acct, k) in enumerate(combos, 2):
            p.append([acct, k,
                      f'=COUNTIFS({rng},A{i},{krng},B{i},{res},"<>open")',
                      f'=COUNTIFS({rng},A{i},{krng},B{i},{res},"W")',
                      f'=COUNTIFS({rng},A{i},{krng},B{i},{res},"L")',
                      f'=IF(D{i}+E{i}=0,"",D{i}/(D{i}+E{i}))',
                      f'=SUMIFS(Plays!$L$2:$L${n2},{rng},A{i},{krng},B{i})',
                      f'=SUMIFS(Plays!$K$2:$K${n2},{rng},A{i},{krng},B{i})',
                      f'=IF(C{i}=0,"",G{i}/C{i})'])
            for col, fmt in (("F", "0%"), ("G", "0.00;[Red]-0.00"), ("H", "#,##0;[Red]-#,##0"),
                             ("I", "0.00;[Red]-0.00")):
                p[f"{col}{i}"].number_format = fmt
        p.auto_filter.ref = f"A1:I{max(len(combos) + 1, 1)}"

    pivot("By League", "League", "C", "league")
    pivot("By Type", "Type", "D", "mtype")
    wb.save(path)
    return path
